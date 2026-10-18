"""Build docs/submission/Refill-Citations.xlsx: every claim Refill makes, with its source or evidence.

Run from this folder:  python citations.py
Then recalculate the formulas (LibreOffice or Excel opens it fine; the xlsx skill's recalc.py also works).
Sources were opened and read on 18 October 2026. Re-check them before reusing this file later.
"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from pathlib import Path

OUT = Path(__file__).resolve().parents[2] / "submission" / "Refill-Citations.xlsx"
ACCESSED = "18 Oct 2026"

FONT = "Arial"
INK, PAPER, MUTED = "121214", "F4F2ED", "6B6B70"
HEAD_FILL = PatternFill("solid", fgColor="121214")
BAND_FILL = PatternFill("solid", fgColor="F7F6F2")
INPUT_BLUE = "0000FF"
STATUS_FILL = {  # Refill's colours by meaning: lime = fine, amber = attention, coral = fix
    "Verified": "DDEFC4",
    "Verified with caveat": "FBE7C2",
    "Needs correction": "F8D0C8",
    "Not verified": "E4E4E7",
    "Passed (observed)": "DDEFC4",
    "Demo data (synthetic)": "FBE7C2",
    "Hypothesis": "E4E4E7",
    "Matches": "DDEFC4",
}
thin = Side(style="thin", color="D9D7D0")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")


def cf_fill(color):
    # Conditional formats read the background colour, not the foreground one.
    return PatternFill("solid", fgColor=color, bgColor=color)

DECK = "Final deck (dark + light)"
BRIEF = "Financial Projections Brief"
MBA = "MBA Master document"
FILM = "Product film (youtu.be/5BfsKofkIHw)"
GUIDE = "MBA guide (public/guide/mba.html)"
CASE_URL = "https://drive.google.com/file/d/1lHi0dQghhwh15UM8cXO0HWz_uvdSb3Xl/view"

# ID, claim, where, topic, source title, publisher/author, url, published, what the source says, status, note
CITATIONS = [
    ("C01", "The brief: fill genuinely available appointment slots and reduce missed sessions without overbooking; show cancellations and waiting-list matches, record consent, rebook, prevent duplicate slot allocation.",
     f"{DECK} slides 1-2; {FILM} lines 1-2; the whole product", "Case requirement",
     "Use Case Book v3, C29 Appointment Recovery Desk, p. 32", "Fynd x NMIMS Hyderabad (case bank)", CASE_URL, "16 Oct 2026",
     "C29 card: user and end goal, engineering build, MBA assignment, growth experiment, product change, starter data.", "Verified",
     "Organiser document, shared on the event Drive."),
    ("C02", "Three scored case checks: a 60-minute request cannot take a 30-minute slot (10 pts); two competing requests give one confirmed booking and one waitlist entry (10 pts); a customer without contact consent is excluded (5 pts).",
     f"{DECK} slide 4; {FILM} line 20; the Lab (#/lab)", "Case requirement",
     "Use Case Book v3, C29, p. 32, 'Case outcome checks for engineering'", "Fynd x NMIMS Hyderabad (case bank)", CASE_URL, "16 Oct 2026",
     "Lists the three checks with 10, 10 and 5 points.", "Verified", "Our proof for each check is on the Product evidence sheet (E01-E03)."),
    ("C03", "Market shocks we prepared for: late cancellation, staff absence, a different service duration, a customer withdrawing consent.",
     f"{DECK} slides 4-5; docs/SHOCKS.md; {MBA} 'Market shocks'", "Case requirement",
     "Use Case Book v3, C29, p. 32, 'Challenge and edge cases'", "Fynd x NMIMS Hyderabad (case bank)", CASE_URL, "16 Oct 2026",
     "Names the same four edge cases.", "Verified", ""),
    ("C04", "The negotiated product change is 'add service-duration and staff constraints before outreach'; the growth experiment uses five consenting role players; pricing is a per-location subscription.",
     f"{DECK} slides 3, 4, 6, 7; {MBA} sections 8, 11, 12", "Case requirement",
     "Use Case Book v3, C29, p. 32, MBA assignment and product change", "Fynd x NMIMS Hyderabad (case bank)", CASE_URL, "16 Oct 2026",
     "'Test per-location subscription', 'Test opt-in waitlist messages with five consenting role players', 'Add service-duration and staff constraints before outreach'.", "Verified", ""),
    ("C05", "The 30-day pilot on one calendar measures filled cancellations, no-show rate and staff correction effort.",
     f"{DECK} slide 7; {FILM} line 22; {MBA} section 13", "Case requirement",
     "Use Case Book v3, C29, p. 32, 'Path to a real pilot'", "Fynd x NMIMS Hyderabad (case bank)", CASE_URL, "16 Oct 2026",
     "'Pilot one calendar; measure filled cancellations, no-show rate and staff correction effort.'", "Verified", ""),
    ("C06", "India's beauty-salon market is about USD 12.6 billion (2025).",
     f"{BRIEF} section 08", "Market size",
     "India Beauty Salon Market Size, Share & Forecast, 2026-2031", "Ken Research",
     "https://www.kenresearch.com/industry-reports/india-beauty-salon-market.md", "2026 (base year 2025)",
     "USD 12,565 million service revenue in 2025; forecast USD 19,773 million by 2031.", "Verified with caveat",
     "Research-firm estimate; method not public. Present as indicative."),
    ("C07", "About 3,00,000 salon players in India.",
     f"{BRIEF} section 08", "Market size",
     "India Beauty Salon Market Size, Share & Forecast, 2026-2031", "Ken Research",
     "https://www.kenresearch.com/industry-reports/india-beauty-salon-market.md", "2026 (base year 2025)",
     "'Total Number of Players: 300,000'; 'approximately 300,000 operators (2025, India)'.", "Verified with caveat",
     "Research-firm estimate."),
    ("C08", "Organised salon outlets in India are above 13,000.",
     f"{BRIEF} section 08", "Market size",
     "India Salon Market Outlook 2032", "Nexdigm",
     "https://www.nexdigm.com/market-research/report-store/india-salon-market/", "2026",
     "'organized salons crossed 13,000 outlets nationwide'.", "Verified with caveat",
     "Nexdigm attributes the figure to 'public industry trackers'."),
    ("C09", "Lakme Salon runs 450+ outlets.",
     f"{BRIEF} section 08", "Market size",
     "India Beauty Salon Market Size, Share & Forecast, 2026-2031", "Ken Research",
     "https://www.kenresearch.com/industry-reports/india-beauty-salon-market.md", "2026",
     "'Lakme Salon operates more than 450 salons across 190 cities (2025, India)'.", "Needs correction",
     "The number is right, but the brief cites Inventiva for it. Cite Ken Research instead."),
    ("C10", "Naturals runs 800+ outlets.",
     f"{BRIEF} section 08", "Market size",
     "India Beauty Salon Market Size, Share & Forecast, 2026-2031", "Ken Research",
     "https://www.kenresearch.com/industry-reports/india-beauty-salon-market.md", "2026",
     "'Naturals operates more than 800 outlets (2025, India)'.", "Needs correction",
     "The brief cites Inventiva, whose page says 650+ (inventiva.co.in/?p=412093). Cite Ken Research, or say '650-800+ depending on source'."),
    ("C11", "Hair and beauty businesses lose almost 7% of monthly revenue to cancellations.",
     f"{BRIEF} section 06", "Customer pain",
     "The Hidden Cost of Cancellations: How No-Shows Are Reshaping the Beauty Industry", "Fresha",
     "https://www.fresha.com/blog/cancellations-cost-study", "2 Feb 2026 (updated 4 Feb 2026)",
     "'In hair and beauty, average revenue lost to cancellations sits at almost 7% a month'.", "Verified with caveat",
     "A survey of UK businesses by a booking-software vendor. Say 'UK survey' when quoting it."),
    ("C12", "Small-business software commonly loses 3-7% of customers a month.",
     f"{BRIEF} section 05", "SaaS benchmark",
     "SaaS Churn Rate: What It Is, How to Calculate It, and Benchmarks", "Kalungi (Fadi George)",
     "https://kalungi.com/blog/saas-churn-rate-benchmarks", "5 May 2025",
     "'For small to medium-sized businesses (SMBs), which typically bill monthly, you'll see a churn rate between 3 and 7%.'", "Verified with caveat",
     "Vendor blog benchmark."),
    ("C13", "The same 3-7% monthly churn range, second source.",
     f"{BRIEF} section 09 (Sources)", "SaaS benchmark",
     "SaaS Churn Rate Benchmarks: What's Actually Normal in 2026", "UserJot",
     "https://userjot.com/blog/saas-churn-rate-benchmarks", "2026",
     "Says 'SMB-focused products: 5-7% annual is realistic' and 'anything over 2-3% monthly is a problem'.", "Needs correction",
     "UserJot does not support 3-7% a month. Keep it only for row C14."),
    ("C14", "Much of software churn happens in the first three months.",
     f"{BRIEF} section 05", "SaaS benchmark",
     "SaaS Churn Rate Benchmarks: What's Actually Normal in 2026", "UserJot",
     "https://userjot.com/blog/saas-churn-rate-benchmarks", "2026",
     "'The majority of churn happens in the first 90 days.'", "Verified with caveat", "Vendor blog."),
    ("C15", "Payment gateway fees are about 2% of the ₹699 subscription.",
     f"{BRIEF} section 04; {MBA} section 9", "Cost to serve",
     "Razorpay pricing", "Razorpay", "https://razorpay.com/pricing/", "Live page",
     "Standard plan fee 2% (new accounts currently get 0% for the first 90 days).", "Verified", ""),
    ("C16", "A WhatsApp utility message costs about ₹0.12-0.15.",
     f"{BRIEF} section 04; {MBA} section 9", "Cost to serve",
     "WhatsApp Business API pricing in India (Meta rate card summary)", "Whautomate",
     "https://whautomate.com/whatsapp-business-api-pricing-india", "Rate card effective 1 Oct 2026",
     "Utility messages ₹0.115 each; from 1 Oct 2026 also charged inside the service window. Provider fees can add 10-30%.", "Verified with caveat",
     "Secondary summary; Meta's own pricing page blocks automated reading. Official page: developers.facebook.com/documentation/business-messaging/whatsapp/pricing"),
    ("C17", "An SMS costs about ₹0.15-0.25.",
     f"{BRIEF} section 04; {MBA} section 9", "Cost to serve",
     "SMS pricing, India to India", "MSG91", "https://msg91.com/in/pricing/sms", "Live page",
     "₹0.16 to ₹0.25 per SMS depending on volume, plus 18% GST.", "Verified", "Lowest published rate is ₹0.16, not ₹0.15."),
    ("C18", "SMS needs DLT registration and approved templates.",
     f"{MBA} section 15 (risks)", "Regulation",
     "Telecom Commercial Communications Customer Preference Regulations, 2018", "Telecom Regulatory Authority of India (TRAI)",
     "https://www.trai.gov.in/sites/default/files/2024-09/RegulationUcc19072018.pdf", "19 Jul 2018",
     "Defines the Consent Register as a Distributed Ledger and requires content templates to be registered for transactional and service messages.", "Verified", ""),
    ("C19", "WhatsApp business messages need approved templates.",
     f"{MBA} section 15 (risks)", "Regulation / platform rule",
     "Template fundamentals, WhatsApp Business Platform", "Meta for Developers",
     "https://developers.facebook.com/documentation/business-messaging/whatsapp/templates/overview", "Live page",
     "Templates must have APPROVED status before they can be sent; outside the 24-hour window only templates can be sent.", "Verified with caveat",
     "Confirmed through search results that quote this page; the page itself blocks automated reading."),
    ("C20", "India's Digital Personal Data Protection Act, 2023 requires clear consent.",
     f"{MBA} section 15; consent-first design", "Regulation",
     "The Digital Personal Data Protection Bill, 2023 (bill summary)", "PRS Legislative Research",
     "https://prsindia.org/billtrack/digital-personal-data-protection-bill-2023", "Aug 2023",
     "'Personal data may be processed only for a lawful purpose after obtaining the consent of the individual. A notice must be given before seeking consent.'", "Verified",
     "Pinpoint in the Act: section 5 (notice) and section 6 (consent)."),
    ("C21", "Supabase Pro costs about US$25 a month.",
     f"{MBA} section 9; {BRIEF} section 04", "Cost to serve",
     "Supabase pricing", "Supabase", "https://supabase.com/pricing", "Live page",
     "'Pro ... From $25 / month'.", "Verified", ""),
    ("C22", "Vercel Pro costs about US$20 a month.",
     f"{MBA} section 9; {BRIEF} section 04", "Cost to serve",
     "Vercel pricing", "Vercel", "https://vercel.com/pricing", "Live page",
     "'Pro $20 /mo' with $20 of included usage credit.", "Verified", ""),
    ("C23", "Hosting costs are converted at about ₹85 per US$.",
     f"{MBA} section 9", "Cost to serve",
     "USD exchange rates (daily)", "ExchangeRate-API", "https://open.er-api.com/v6/latest/USD", "18 Oct 2026",
     "1 USD = ₹95.08 on 18 Oct 2026.", "Needs correction",
     "Use about ₹95. The hosting share for 50 locations rises from about ₹77 to about ₹86 per location a month; the ₹100 COGS still covers it."),
    ("C24", "Mangomint costs $165-$375 a month by team size.",
     f"{MBA} section 7 (competitor table)", "Competitor",
     "Mangomint pricing", "Mangomint", "https://www.mangomint.com/pricing/", "Live page",
     "'Monthly price is $120 per location and $10 per user', with paid add-ons.", "Needs correction",
     "The MBA figure is out of date. Use '$120 per location + $10 per user'."),
    ("C25", "Mangomint handles waitlists with express booking, virtual waiting rooms and automated client flows.",
     f"{MBA} section 7", "Competitor",
     "Mangomint features", "Mangomint", "https://www.mangomint.com/features/", "Live page",
     "Feature list includes Express Booking, Virtual Waiting Room and Automated Flows.", "Verified", ""),
    ("C26", "Vagaro starts at $30 a month for one calendar, more per staff calendar, with waitlist tracking.",
     f"{MBA} section 7", "Competitor",
     "Vagaro pricing", "Vagaro", "https://www.vagaro.com/pro/pricing", "Live page",
     "$30/month for one location ($23.99 introductory); each extra bookable calendar $10, up to seven; 'Customer waitlist options' included.", "Verified", ""),
    ("C27", "Fresha has a low base price plus transaction and messaging fees, with waitlists.",
     f"{MBA} section 7; {DECK} slide 4", "Competitor",
     "Fresha pricing (India)", "Fresha", "https://www.fresha.com/pricing", "Live page",
     "Free subscription for one person; Team plan charges per bookable team member; texts ₹13.50 and WhatsApp ₹1.60-5.50 after 20 free a month; 'Waitlists' listed as a feature.", "Verified with caveat",
     "Payment-processing fees are not shown on this page."),
    ("C28", "Zenoti offers a waitlist.",
     f"{DECK} slide 4 (gap vs competitors)", "Competitor",
     "Zenoti salon software", "Zenoti", "https://www.zenoti.com/salon-software", "Live page",
     "Feature list includes 'Automated Waitlist Management'.", "Verified", ""),
    ("C29", "Square offers a waitlist.",
     f"{DECK} slide 4 (gap vs competitors)", "Competitor",
     "Square Appointments features", "Square", "https://squareup.com/us/en/appointments/features", "Live page",
     "'Create a waitlist for popular services and fill gaps in your schedule due to last-minute cancellations or reschedules.'", "Verified", ""),
    ("C30", "Zenoti has its headquarters in Hyderabad.",
     f"{MBA} section 7", "Competitor",
     "Zenoti company page and contact page", "Zenoti", "https://www.zenoti.com/contact-us", "Live page",
     "HQ: 15900 SE Eastgate Way, Bellevue WA, USA. The company page says Zenoti began as ManageMySpa in Hyderabad in 2010.", "Needs correction",
     "Say 'founded in Hyderabad, headquartered in Bellevue, USA'."),
    ("C31", "TutorBird costs $16.95 a month solo plus $4.95 per extra tutor, with student waitlists.",
     f"{MBA} section 7", "Competitor",
     "TutorBird pricing", "TutorBird", "https://www.tutorbird.com/pricing/", "Live page",
     "'$16.95 /month + $4.95 per each additional tutor or staff member'; 'Track your waitlist & leads'.", "Verified with caveat",
     "'Cancelled lessons become make-up slots' is not stated on the pages checked."),
    ("C32", "Calendly and Acuity cost $10-$30+ per seat a month.",
     f"{MBA} section 7", "Competitor",
     "Calendly pricing; Acuity Scheduling plans", "Calendly; Squarespace (Acuity)",
     "https://calendly.com/pricing", "Live page",
     "Calendly Standard $10/seat/month, Teams $16/seat/month; Acuity Starter $16-$20/month (acuityscheduling.com/signup.php).", "Verified with caveat",
     "The top of the range depends on plan; enterprise plans are custom."),
    ("C33", "\"Rich people value time\": higher-income clients feel short of time, so a relevant opt-in offer of an earlier slot is credible.",
     f"{DECK} slide 3 (psychographic insight)", "Customer insight",
     "Stressed Out on Four Continents: Time Crunch or Yuppie Kvetch? (NBER Working Paper 10186)", "Daniel S. Hamermesh and Jungmin Lee",
     "https://www.nber.org/papers/w10186", "2003; Review of Economics and Statistics, 2007",
     "'Time stress will be more prevalent in households with higher incomes'; corroborated in Australia, Canada, Germany, Korea and the US.", "Verified with caveat",
     "Research on time stress in general, not on salon clients."),
    ("C34", "People who pay to save time are happier, which supports selling convenience to time-poor clients.",
     f"{DECK} slide 3 (supporting reading)", "Customer insight",
     "Buying time promotes happiness, PNAS 114(32): 8523-8527", "Ashley V. Whillans, Elizabeth W. Dunn, Paul Smeets, Rene Bekkers, Michael I. Norton",
     "https://doi.org/10.1073/pnas.1706541114", "Jul 2017",
     "Spending on time-saving services was linked to greater life satisfaction across four countries; a field experiment found the same.", "Verified with caveat",
     "Supporting evidence only; the deck quotes a team insight, not this paper."),
    ("C35", "A Booksy survey of US salons found only about 2% of clients cancel without any warning.",
     GUIDE, "Customer pain",
     "No-Shows: Does Your City Have a High Rate?", "Salon Today (reporting a Booksy survey)",
     "https://www.salontoday.com/1076870/no-shows-does-your-city-have-a-high-rate", "23 Sep 2021",
     "'at worst over time only about 2 percent of clients don't even give the salon a heads-up'; 15 US cities, 0.98% to 2.37%.", "Verified with caveat",
     "2021 US data, reported second-hand."),
    ("C36", "Treatwell (UK) reports 2-3% no-shows per treatment, and says salons requiring prepayment halved theirs.",
     GUIDE, "Customer pain",
     "Beauty salon no-shows: the UK locations with the highest rates", "Scratch Magazine (reporting Treatwell research)",
     "https://www.scratchmagazine.co.uk/news/beauty-salon-no-shows-uk-locations-highest-rates/", "Mar 2025",
     "Highest no-show treatments run 2.69% to 3.14%; '17% of salons now requiring pre-payment, which has reduced no-shows by half'.", "Verified with caveat",
     "2-3% is the range for the most-missed treatments, not every treatment."),
]

# ID, claim, where, evidence, observed result, status
EVIDENCE = [
    ("E01", "A 60-minute request never takes a 30-minute slot.", f"{DECK} slide 4; {FILM} line 20",
     "Fit check in src/engine.ts and in the database functions (supabase/migrations); scenario C02 in tests/core-scenarios.json; Lab case check on #/lab.",
     "npm test 51/51 and npm run dbtest 32/32 passed on 18 Oct 2026.", "Passed (observed)"),
    ("E02", "Two requests for one slot give one confirmed booking; the other client keeps their waitlist place.", f"{DECK} slides 4, 7; {FILM} lines 13-14",
     "Booking decided inside the database under a row lock plus a unique index; scenarios C04, C09, C13, C14; npm run race (8 simultaneous taps).",
     "npm run race: 3 of 3 runs gave one booking and 7 clients kept their place; a repeated tap returned the same booking (18 Oct 2026).", "Passed (observed)"),
    ("E03", "A customer without contact consent is excluded from outreach, with the reason logged.", f"{DECK} slide 4; {FILM} line 9",
     "Scenarios C05, C06, C12; reason code NO_CONSENT in the evidence log.", "npm test and npm run dbtest passed on 18 Oct 2026.", "Passed (observed)"),
    ("E04", "Consent is checked when offers go out and again at booking; STOP withdraws every open offer at once.", f"{DECK} slide 4 ('TESTS C05 · C06 · C12'); {FILM} line 12",
     "Scenarios C05, C06, C12; race check 'STOP after the offer: booking refused, slot stays free'.", "Passed in npm run race and npm run dbtest on 18 Oct 2026.", "Passed (observed)"),
    ("E05", "Offers are held 15 minutes; at most 3 offers per client a day; under-2-hour cancellations are flagged late; no offers inside 10 minutes of the start; mistakes can be undone.", f"{DECK} slides 4-5; {FILM} line 18",
     "Database rules OFFER_EXPIRED, FREQUENCY_CAP, late flag and restore_slot (supabase/migrations 002-003); Owner > Rules page. These are defaults the owner can change.",
     "Covered by npm run dbtest (32/32) on 18 Oct 2026.", "Passed (observed)"),
    ("E06", "When staff can't come in, their offers are withdrawn, their gaps held and their clients flagged; the desk moves them or sends a pick-a-time link.", f"{DECK} slides 4-5; {FILM} line 15",
     "Staff absence and moves (supabase/migrations 004-005); scenario C20; npm run flowtest.", "npm run flowtest 14/14 passed on 18 Oct 2026.", "Passed (observed)"),
    ("E07", "Every cancellation, offer (sent or skipped, with reason), booking, expiry, move and undo is logged and can be exported as CSV or JSON.", f"{DECK} slide 5; {FILM} line 18",
     "Append-only events table (migration 003); Activity export in the app.", "Visible in the live app; covered by dbtest.", "Passed (observed)"),
    ("E08", "The rules live in the database, not just in the app.", f"{FILM} lines 19-20; {DECK} slide 7",
     "Functions-only API: tables have no direct client access; every action is a database function (supabase/migrations).", "dbtest calls the live database directly and the rules still hold (32/32).", "Passed (observed)"),
    ("E09", "51 automated tests, 28 live database checks and 14 cross-role flow checks all pass.", f"{MBA} sections 12, 16",
     "npm test, npm run dbtest, npm run flowtest.", "51/51, 32/32 (four checks added on 18 Oct for absence cover) and 14/14 on 18 Oct 2026.", "Passed (observed)"),
    ("E10", "Demo day: 17 people waiting; 6 fit and 11 are skipped (4 no consent, 2 wrong skill, 2 need more time, 3 not free).", f"{DECK} slide 4 ('6 FIT · 11 SKIPPED'); {MBA} section 5",
     "Demo seed for the made-up Kesar Salon (synthetic data, data/salon/*.csv).", "Reproduced on the live demo after 'Reset demo day'.", "Demo data (synthetic)"),
    ("E11", "One late cancellation put ₹1,600 at risk and Refill recovered it in two taps.", f"{MBA} sections 5, 16; {FILM} line 16",
     "Harsh's 12:00 hair spa in the demo seed.", "Reproduced on the live demo.", "Demo data (synthetic)"),
    ("E12", "Everything shown runs on synthetic data; no live revenue is claimed.", f"{FILM} line 23; {DECK}; {BRIEF}",
     "The case book allows synthetic starter data (Use Case Book p. 1); all figures are tagged 'demo data'.", "Checked in the deck, film and brief.", "Passed (observed)"),
    ("E13", "When someone is marked out, their coming appointments move automatically to a teammate who is in, does that service and is free at that time; the rest stay flagged.", "Live app (owner rule 'Cover absences automatically'); README",
     "Absence planner cover_absence_plan (supabase/migrations/012_absence_cover.sql); db-test check 12c.", "npm run dbtest 32/32 on 18 Oct 2026; on the demo day, marking Arjun out moves Manav's 16:30 to Riya and flags three clients.", "Passed (observed)"),
    ("E14", "The owner is notified of every absence with what was covered, what is left, the person's tasks and who is in; tasks can be handed over or marked missed.", "Live app (Owner overview, bell); README",
     "notices table, reassign_task and read_notice (migration 012); db-test check 12c.", "npm run dbtest 32/32 on 18 Oct 2026; checked in the browser on the demo day.", "Passed (observed)"),
]

# ID, claim, where, basis, how we test it
ASSUMPTIONS = [
    ("A01", "Price ₹699 per location per month is the anchor; test ₹499, ₹699 and ₹999.", f"{DECK} slides 2, 6; {BRIEF}",
     "MBA team hypothesis. Only a paid-pilot decision counts as validation.", "Owner discovery interviews and the paid pilot."),
    ("A02", "MBA Master section 8 tests ₹999 / ₹1,499 / ₹1,999 (H1: ₹1,499).", f"{MBA} section 8",
     "Earlier MBA hypothesis.", "Inconsistent with the deck's ₹699 anchor. Align the document with the deck before submitting."),
    ("A03", "Cost to serve is ₹100 per location per month (hosting, messages, payment fees).", f"{DECK} slide 7; {BRIEF} section 04",
     "Built from rows C15-C17 and C21-C22; see the 'Our numbers' sheet.", "Real invoices once paid plans start."),
    ("A04", "Average active locations: 50, 250 and 800 in years 1-3.", f"{DECK} slide 7; {BRIEF} section 03",
     "Bottom-up: about 8, 25 and 65 net new locations a month.", "The 30-day pilot and the first sales months."),
    ("A05", "Operating costs: ₹15 L, ₹27 L and ₹49 L in years 1-3.", f"{DECK} slide 7; {BRIEF} section 05",
     "Illustrative split: team, sales, marketing, tools, compliance.", "Budget review once the team is hired."),
    ("A06", "At ₹500 contribution per service, about 2 recovered bookings a month cover ₹699.", f"{DECK} slide 6",
     "₹500 contribution per service is an assumption.", "Owner baseline in pilot week 3."),
    ("A07", "Worked example: a salon taking ₹3,00,000 a month loses about ₹21,000 to cancellations.", f"{BRIEF} section 06",
     "Revenue is an assumption; the 7% comes from row C11 (UK survey).", "Owner baseline in pilot week 3."),
    ("A08", "Target customer: owner-operated, one location, 3-15 staff, Hyderabad first.", f"{DECK} slide 3",
     "Segmentation hypothesis.", "Owner interviews and the pilot."),
    ("A09", "Competitors do not expose auditable consent state, duration-mismatch rejection or slot-contention proof to the operator.", f"{DECK} slide 4",
     "Our reading of competitor product pages; the deck itself says 'Validate by demo before claiming'.", "Hands-on demos of Fresha, Zenoti and Square."),
    ("A10", "Calendly and Acuity have weak or no automatic waitlist backfill.", f"{MBA} section 7",
     "Our reading of their product pages; absence is hard to prove.", "Hands-on trial of both."),
    ("A11", "Role-player success: at least 3 of 5 respond within the 15-minute hold; at least 4 of 5 rate the offer 4 or 5 for clarity.", f"{MBA} section 11",
     "Targets we set before the test.", "The five-role-player experiment."),
    ("A12", "Pilot success: filled cancellations rise clearly above baseline (for example, from under 20% to over 50%) with no rise in no-shows.", f"{MBA} section 13",
     "Example threshold; agree the exact numbers with the owner before the pilot.", "The 30-day pilot."),
    ("A13", "Messaging costs about ₹10-25 per location per month (about 90 messages).", f"{MBA} section 14",
     "3 offers x 30 cancellations; rates from rows C16-C17; see 'Our numbers'.", "Pilot weeks 2-3."),
    ("A14", "Set-up takes 30-60 minutes; desk training takes 15 minutes.", f"{MBA} sections 9, 13",
     "Our estimate from demo walk-throughs.", "Time it in the pilot."),
    ("A15", "A 20% salon no-show rate is widely repeated.", GUIDE,
     "Flagged in the guide as unverified; it is not used in the deck.", "Do not quote it."),
]


def style_table(ws, header, rows, widths, status_col=None):
    ws.append(header)
    for c in ws[1]:
        c.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = BOX
    for r in rows:
        ws.append(list(r))
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name=FONT, size=10, color=INK)
            c.alignment = WRAP
            c.border = BOX
        row[0].font = Font(name=FONT, size=10, bold=True, color=INK)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 30
    if status_col:
        last = ws.max_row
        rng = f"{status_col}2:{status_col}{last}"
        for name, color in STATUS_FILL.items():
            ws.conditional_formatting.add(rng, FormulaRule(formula=[f'${status_col}2="{name}"'], fill=cf_fill(color)))


def build():
    wb = Workbook()

    # Read me
    ws = wb.active
    ws.title = "Read me"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 90
    lines = [
        ("Refill: citations and evidence", None),
        ("Team C29 · Fynd x NMIMS Hyderabad Product & Business Simulation · Case C29 Appointment Recovery Desk", None),
        ("", None),
        ("What this file is", "Every claim Refill makes in the deck, the documents, the film and the guides, with where it came from. Outside facts are on 'Citations'. Claims we prove with our own tests are on 'Product evidence'. Our own estimates are on 'Assumptions'. 'Our numbers' rebuilds the deck's financial table with live formulas."),
        ("Sources checked on", f"{ACCESSED}. Each link was opened and read; the 'What the source says' column quotes or paraphrases what it said that day."),
        ("Live app", "https://fillthechair.vercel.app"),
        ("Repository", "https://github.com/sirgreaseball/refill"),
        ("", None),
        ("Status key", None),
        ("Verified", "The source says what we say."),
        ("Verified with caveat", "The source supports the claim, with a limit worth saying aloud (a vendor survey, another country, a secondary report)."),
        ("Needs correction", "The source says something different or the citation points to the wrong place. The note gives the fix."),
        ("Not verified", "No source found that confirms it."),
        ("Passed (observed)", "Our own test or check, run on the live system on the date shown."),
        ("Demo data (synthetic)", "A number from the made-up demo salon. Never presented as real revenue."),
        ("Hypothesis", "Our assumption, with the test that will confirm or reject it."),
        ("", None),
        ("Counts", None),
    ]
    for b, c in lines:
        ws.append(["", b, c])
    r0 = ws.max_row
    counts = [
        ("Outside sources (Citations)", "=COUNTA(Citations!A2:A200)"),
        ("  Verified", '=COUNTIF(Citations!J2:J200,"Verified")'),
        ("  Verified with caveat", '=COUNTIF(Citations!J2:J200,"Verified with caveat")'),
        ("  Needs correction", '=COUNTIF(Citations!J2:J200,"Needs correction")'),
        ("  Not verified", '=COUNTIF(Citations!J2:J200,"Not verified")'),
        ("Product claims with evidence", "=COUNTA('Product evidence'!A2:A200)"),
        ("Assumptions", "=COUNTA(Assumptions!A2:A200)"),
        ("Financial table cells that match the deck", "=COUNTIF('Our numbers'!F16:F40,\"Matches\")&\" of \"&(COUNTIF('Our numbers'!F16:F40,\"Matches\")+COUNTIF('Our numbers'!F16:F40,\"Differs\"))"),
    ]
    for b, f in counts:
        ws.append(["", b, f])
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        for c in row:
            c.font = Font(name=FONT, size=10, color=INK)
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws["B1"].font = Font(name=FONT, size=18, bold=True, color=INK)
    ws["B2"].font = Font(name=FONT, size=10, color=MUTED)
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, 2).value in ("What this file is", "Sources checked on", "Live app", "Repository", "Status key", "Counts"):
            ws.cell(r, 2).font = Font(name=FONT, size=10, bold=True, color=INK)
        name = ws.cell(r, 2).value
        if name in STATUS_FILL and r < r0:
            ws.cell(r, 2).fill = PatternFill("solid", fgColor=STATUS_FILL[name])
    for r in range(r0 + 1, ws.max_row + 1):
        ws.cell(r, 3).alignment = Alignment(horizontal="left", vertical="top")
    ws["C6"].hyperlink = "https://fillthechair.vercel.app"
    ws["C7"].hyperlink = "https://github.com/sirgreaseball/refill"
    for ref in ("C6", "C7"):
        ws[ref].font = Font(name=FONT, size=10, color="1F5FBF", underline="single")

    # Citations
    ws = wb.create_sheet("Citations")
    header = ["ID", "Claim as we made it", "Where we made it", "Topic", "Source", "Publisher / author", "Link",
              "Published", "What the source says", "Status", "Note / fix", "Accessed"]
    rows = [r + (ACCESSED,) for r in CITATIONS]
    style_table(ws, header, rows, [6, 44, 26, 14, 34, 22, 34, 14, 46, 15, 36, 11], status_col="J")
    for r in range(2, ws.max_row + 1):
        url = ws.cell(r, 7).value
        for col in (5, 7):
            ws.cell(r, col).hyperlink = url
            ws.cell(r, col).font = Font(name=FONT, size=10, color="1F5FBF", underline="single")
    dv = DataValidation(type="list", formula1='"Verified,Verified with caveat,Needs correction,Not verified"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"J2:J{ws.max_row}")

    # Product evidence
    ws = wb.create_sheet("Product evidence")
    style_table(ws, ["ID", "Claim", "Where we made it", "Evidence (file, test or page)", "Observed result", "Status"],
                EVIDENCE, [6, 48, 30, 52, 40, 18], status_col="F")

    # Assumptions
    ws = wb.create_sheet("Assumptions")
    rows = [r + ("Hypothesis",) for r in ASSUMPTIONS]
    style_table(ws, ["ID", "Assumption", "Where we used it", "Basis", "How we test it", "Status"],
                rows, [6, 52, 30, 46, 40, 12], status_col="F")

    # Our numbers: the deck's financial table rebuilt with formulas
    ws = wb.create_sheet("Our numbers")
    ws.sheet_view.showGridLines = False
    widths = [40, 14, 14, 14, 16, 12]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    blue = Font(name=FONT, size=10, color=INPUT_BLUE)
    plain = Font(name=FONT, size=10, color=INK)
    bold = Font(name=FONT, size=10, bold=True, color=INK)
    ws["A1"] = "The financial projections, rebuilt with formulas"
    ws["A1"].font = Font(name=FONT, size=14, bold=True, color=INK)
    ws["A2"] = "Blue cells are inputs (hypotheses, see 'Assumptions'). Black cells are formulas. Column E is what the deck shows; column F checks it."
    ws["A2"].font = Font(name=FONT, size=9, color=MUTED)
    inputs = [
        ("Price per location per month (₹)", 699, "A01"),
        ("Cost to serve per location per month (₹)", 100, "A03"),
        ("Months in a year", 12, ""),
        ("Rupees in one lakh", 100000, ""),
        ("Contribution per service (₹)", 500, "A06"),
        ("Salon revenue in the worked example (₹ / month)", 300000, "A07"),
        ("Revenue lost to cancellations (hair and beauty)", 0.07, "C11"),
        ("All salon players in India", 300000, "C07"),
        ("Organised salon outlets", 13000, "C08"),
    ]
    ws.append([])
    ws.append(["Inputs", "Value", "Row ref"])
    for c in ws[4]:
        c.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
    for label, val, ref in inputs:
        ws.append([label, val, ref])
        ws.cell(ws.max_row, 1).font = plain
        ws.cell(ws.max_row, 2).font = blue
        ws.cell(ws.max_row, 3).font = Font(name=FONT, size=9, color=MUTED)
    # B5 price, B6 cost, B7 months, B8 lakh, B9 contribution, B10 salon revenue, B11 loss %, B12 players, B13 organised
    ws["B11"].number_format = "0.0%"
    for ref in ("B5", "B6", "B9", "B10", "B12", "B13"):
        ws[ref].number_format = "#,##0"

    r = 15
    ws.cell(r, 1, "Projection (₹ lakh)")
    for i, h in enumerate(["Year 1", "Year 2", "Year 3", "Deck (Y1 | Y2 | Y3)", "Check"], start=2):
        ws.cell(r, i, h)
    for c in ws[r]:
        c.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
    table = [
        ("Average active locations", [50, 250, 800], "50 | 250 | 800", True),
        ("Revenue", ["=B16*$B$5*$B$7/$B$8", "=C16*$B$5*$B$7/$B$8", "=D16*$B$5*$B$7/$B$8"], "4.19 | 20.97 | 67.10", False),
        ("Variable COGS", ["=B16*$B$6*$B$7/$B$8", "=C16*$B$6*$B$7/$B$8", "=D16*$B$6*$B$7/$B$8"], "0.60 | 3.00 | 9.60", False),
        ("Gross profit", ["=B17-B18", "=C17-C18", "=D17-D18"], "3.59 | 17.97 | 57.50", False),
        ("Total OpEx", [15, 27, 49], "15.00 | 27.00 | 49.00", True),
        ("EBITDA", ["=B19-B20", "=C19-C20", "=D19-D20"], "-11.41 | -9.03 | 8.50", False),
    ]
    for label, vals, deck, is_input in table:
        r += 1
        ws.cell(r, 1, label).font = plain
        for j, v in enumerate(vals):
            c = ws.cell(r, 2 + j, v)
            c.font = blue if is_input else plain
            c.number_format = "#,##0" if label == "Average active locations" else "0.00;-0.00"
        ws.cell(r, 5, deck).font = Font(name=FONT, size=9, color=MUTED)
        dvals = [float(x) for x in deck.split(" | ")]
        checks = "AND(" + ",".join(f"ROUND({ws.cell(r, 2 + j).coordinate},2)={dvals[j]}" for j in range(3)) + ")"
        ws.cell(r, 6, f'=IF({checks},"Matches","Differs")').font = bold

    r += 2
    ws.cell(r, 1, "Unit economics and checks")
    for i, h in enumerate(["Value", "", "", "Deck / brief says", "Check"], start=2):
        ws.cell(r, i, h)
    for c in ws[r]:
        c.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
    checks = [
        ("Contribution per location per month (₹)", "=B5-B6", "599", "=IF(B{r}=599,\"Matches\",\"Differs\")", "#,##0"),
        ("Gross margin", "=(B5-B6)/B5", "about 86%", "=IF(ROUND(B{r},2)=0.86,\"Matches\",\"Differs\")", "0.0%"),
        ("Break-even locations, year 1 OpEx", "=ROUNDUP(B20*$B$8/((B5-B6)*B7),0)", "209", "=IF(B{r}=209,\"Matches\",\"Differs\")", "#,##0"),
        ("Break-even locations, year 3 OpEx", "=ROUNDUP(D20*$B$8/((B5-B6)*B7),0)", "682", "=IF(B{r}=682,\"Matches\",\"Differs\")", "#,##0"),
        ("Break-even at ₹499, year 3 OpEx", "=ROUNDUP(D20*$B$8/((499-B6)*B7),0)", "about 1,000", "=IF(ABS(B{r}-1000)<=50,\"Matches\",\"Differs\")", "#,##0"),
        ("Break-even at ₹999, year 3 OpEx", "=ROUNDUP(D20*$B$8/((999-B6)*B7),0)", "about 450", "=IF(ABS(B{r}-450)<=10,\"Matches\",\"Differs\")", "#,##0"),
        ("Recovered bookings a month to cover the price", "=B5/B9", "about 2", "=IF(ROUNDUP(B{r},0)=2,\"Matches\",\"Differs\")", "0.0"),
        ("Worked example: lost to cancellations (₹ / month)", "=B10*B11", "about ₹21,000", "=IF(ROUND(B{r},0)=21000,\"Matches\",\"Differs\")", "#,##0"),
        ("Recovering one-sixth of that, as a multiple of the price", "=B10*B11/6/B5", "five times", "=IF(ROUND(B{r},0)=5,\"Matches\",\"Differs\")", "0.0\"x\""),
        ("800 locations as a share of all players", "=D16/B12", "under 0.3%", "=IF(B{r}<0.003,\"Matches\",\"Differs\")", "0.00%"),
        ("800 locations as a share of organised outlets", "=D16/B13", "about 6%", "=IF(ROUND(B{r},2)=0.06,\"Matches\",\"Differs\")", "0.0%"),
        ("30% slower adoption, year 1 locations", "=ROUND(B16*0.7,0)", "35", "=IF(B{r}=35,\"Matches\",\"Differs\")", "#,##0"),
        ("30% slower adoption, year 3 locations", "=ROUND(D16*0.7,0)", "560", "=IF(B{r}=560,\"Matches\",\"Differs\")", "#,##0"),
    ]
    for label, f, deck, chk, fmt in checks:
        r += 1
        ws.cell(r, 1, label).font = plain
        c = ws.cell(r, 2, f)
        c.font = plain
        c.number_format = fmt
        ws.cell(r, 5, deck).font = Font(name=FONT, size=9, color=MUTED)
        ws.cell(r, 6, chk.format(r=r)).font = bold
    for row in ws.iter_rows(min_row=15, max_row=r):
        for c in row:
            c.border = BOX
    ws.conditional_formatting.add(f"F16:F{r}", FormulaRule(formula=['$F16="Matches"'], fill=cf_fill(STATUS_FILL["Matches"])))
    ws.conditional_formatting.add(f"F16:F{r}", FormulaRule(formula=['$F16="Differs"'], fill=cf_fill(STATUS_FILL["Needs correction"])))
    ws.freeze_panes = "A5"

    for s in wb.worksheets:
        s.sheet_properties.tabColor = {"Read me": "121214", "Citations": "8BC34A", "Product evidence": "3F6B00",
                                       "Assumptions": "E3A21A", "Our numbers": "E8735A"}[s.title]
        s.page_setup.orientation = "landscape"
        s.page_setup.fitToWidth = 1
        s.page_setup.fitToHeight = 0
        s.sheet_properties.pageSetUpPr.fitToPage = True
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
